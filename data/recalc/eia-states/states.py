#!/usr/bin/env python3
"""Блок 03 «Счёт за свет»: цены EIA на электричество для домохозяйств в штатах PJM.

Собирает три файла в out/:
  03-states.csv         — калькулятор «Выбери штат» (июль 2026 к июлю 2025 и к июлю 2024)
  03-states-annual.csv  — годовая средняя цена для домохозяйств 2019–2025
  claims-03a.csv        — реестр утверждений 03-33 … 03-59

Источники (public domain, открытые файлы без ключей):
  EIA Electric Power Monthly, табл. 5.6.A (июль 2026, выпуск 24.09.2026)
  EIA Electric Sales, Revenue and Average Price 2024, табл. 5A (выпуск 07.10.2025)
  EIA-861M «HS861M 2010-» (помесячно по штатам, 2025–2026 — предварительно)
  EIA-861 «HS861 2010-» (годовые окончательные данные 2010–2024)
  EIA-861 2024, Sales_Ult_Cust_2024.xlsx (коды балансирующих зон — доля PJM по штатам)
  FRED: CPIAUCNS, CUUR0000SEHF01, CUSR0000SEHF01; BLS CPI News Release (июль, август 2026)
  PJM: страница Territory Served

Запуск: python3 states.py            (скачивает недостающие файлы в out/raw/)
        python3 states.py --refresh  (скачивает всё заново)
"""
import csv, html, os, re, subprocess, sys
from collections import defaultdict
import openpyxl

RAW = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.dirname(RAW)
UA = "data-race-research (research@example.org)"

U = {
    "t56a": "https://www.eia.gov/electricity/monthly/xls/table_5_06_a.xlsx",
    "epm": "https://www.eia.gov/electricity/monthly/",
    "enduse": "https://www.eia.gov/electricity/monthly/update/end-use.php",
    "t5a": "https://www.eia.gov/electricity/sales_revenue_price/xls/table_5A.xlsx",
    "esrap": "https://www.eia.gov/electricity/sales_revenue_price/",
    "hs861m": "https://www.eia.gov/electricity/data/state/xls/861m/HS861M%202010-.xlsx",
    "hs861": "https://www.eia.gov/electricity/data/state/xls/861/HS861%202010-.xlsx",
    "f861": "https://www.eia.gov/electricity/data/eia861/zip/f8612024.zip",
    "pjm": "https://www.pjm.com/about-pjm/who-we-are/territory-served",
    "bls07": "https://www.bls.gov/news.release/archives/cpi_08122026.htm",
    "bls08": "https://www.bls.gov/news.release/archives/cpi_09112026.htm",
    "fred": "https://fred.stlouisfed.org/graph/fredgraph.csv?id={}&cosd=2018-01-01",
    "fredser": "https://fred.stlouisfed.org/series/{}",
}
FILES = {
    "t56a_new.xlsx": U["t56a"], "epm_index.html": U["epm"], "epm_update.html": U["enduse"],
    "table_5A.xlsx": U["t5a"], "esrap.html": U["esrap"], "HS861M_2010.xlsx": U["hs861m"],
    "HS861_2010.xlsx": U["hs861"], "f8612024.zip": U["f861"], "pjm_territory.html": U["pjm"],
    "bls_cpi_08122026.htm": U["bls07"], "bls_cpi_09112026.htm": U["bls08"],
}
FRED_IDS = ["CPIAUCNS", "CPIAUCSL", "CUUR0000SEHF01", "CUSR0000SEHF01"]
for s in FRED_IDS:
    FILES[f"fred_{s}.csv"] = U["fred"].format(s)


def fetch(refresh=False):
    for name, url in FILES.items():
        p = os.path.join(RAW, name)
        if os.path.exists(p) and not refresh:
            continue
        subprocess.run(["curl", "-sS", "-L", "-A", UA, "-o", p, url], check=True)
    z = os.path.join(RAW, "Sales_Ult_Cust_2024.xlsx")
    if not os.path.exists(z) or refresh:
        subprocess.run(["unzip", "-o", "-q", os.path.join(RAW, "f8612024.zip"),
                        "Sales_Ult_Cust_2024.xlsx", "-d", RAW], check=True)


def page_text(name):
    t = open(os.path.join(RAW, name), encoding="utf-8", errors="ignore").read()
    t = re.sub(r"<(script|style).*?</\1>", " ", t, flags=re.S)
    t = re.sub(r"<[^>]+>", " ", t)
    return html.unescape(re.sub(r"\s+", " ", t))


# ---------------------------------------------------------------- штаты
# (англ. название, код, рус. название, охват PJM, пояснение; {s} — доля бытовых потребителей в PJM)
STATES = [
    ("Delaware", "DE", "Делавэр", "целиком", "весь штат в PJM ({s}% бытовых потребителей)"),
    ("Illinois", "IL", "Иллинойс", "частично",
     "только зона ComEd (север штата с Чикаго) — ≈{s}% бытовых потребителей; юг (Ameren Illinois) — в MISO. Цена EIA — среднее по всему штату"),
    ("Indiana", "IN", "Индиана", "частично",
     "только зона Indiana Michigan Power (AEP; северо-восток, Форт-Уэйн) и ряд муниципальных сетей — ≈{s}% бытовых потребителей; остальное — MISO. Цена EIA — среднее по всему штату"),
    ("Kentucky", "KY", "Кентукки", "частично",
     "Kentucky Power (восток), Duke Energy Kentucky (север) и кооперативы East Kentucky Power — ≈{s}% бытовых потребителей; LG&E и KU (Луисвилл, центр) — отдельная балансирующая зона, запад — TVA и MISO. Цена EIA — среднее по всему штату"),
    ("Maryland", "MD", "Мэриленд", "целиком", "весь штат в PJM ({s}% бытовых потребителей)"),
    ("Michigan", "MI", "Мичиган", "частично",
     "только юго-запад (Indiana Michigan Power и несколько местных сетей) — ≈{s}% бытовых потребителей; остальное — MISO. Цена EIA — среднее по всему штату"),
    ("New Jersey", "NJ", "Нью-Джерси", "целиком",
     "весь штат в PJM (≈{s}% бытовых потребителей; единичные — в NYISO)"),
    ("North Carolina", "NC", "Северная Каролина", "частично",
     "только северо-восток (Dominion Energy North Carolina и местные кооперативы) — ≈{s}% бытовых потребителей; остальное — Duke Energy (вне RTO) и TVA. Цена EIA — среднее по всему штату"),
    ("Ohio", "OH", "Огайо", "целиком", "весь штат в PJM ({s}% бытовых потребителей)"),
    ("Pennsylvania", "PA", "Пенсильвания", "целиком",
     "практически весь штат в PJM (≈{s}% бытовых потребителей; округ Пайк — NYISO)"),
    ("Tennessee", "TN", "Теннесси", "частично",
     "только северо-восток (Kingsport Power, AEP) — ≈{s}% бытовых потребителей; остальное — TVA. Цена EIA — среднее по всему штату"),
    ("Virginia", "VA", "Вирджиния", "почти целиком",
     "≈{s}% бытовых потребителей в PJM; крайний юго-запад (Бристоль, Kentucky Utilities, Powell Valley) — TVA и LG&E-KU"),
    ("West Virginia", "WV", "Западная Вирджиния", "целиком", "весь штат в PJM ({s}% бытовых потребителей)"),
    ("District of Columbia", "DC", "округ Колумбия", "целиком", "вся территория в PJM ({s}% бытовых потребителей, Pepco)"),
]
US = ("United States", "US", "США", "не PJM", "среднее по США — для сравнения")


def load_t56a():
    ws = openpyxl.load_workbook(os.path.join(RAW, "t56a_new.xlsx"), data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    title = rows[1][0]                      # 'by State, July 2026 and 2025 (Cents per Kilowatthour)'
    hdr = rows[3]
    assert hdr[1] == "July 2026" and hdr[2] == "July 2025", hdr
    d = {}
    for r in rows[4:]:
        if r[0] and isinstance(r[1], (int, float)):
            d[r[0]] = dict(res=r[1], res_prev=r[2], com=r[3], com_prev=r[4], all=r[9], all_prev=r[10])
    return title, d


def load_t5a():
    ws = openpyxl.load_workbook(os.path.join(RAW, "table_5A.xlsx"), data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0].startswith("2024 Average Monthly Bill"), rows[0][0]
    return {r[0]: dict(cust=r[1], kwh=r[2], price=r[3], bill=r[4])
            for r in rows[3:] if r[0] and isinstance(r[2], (int, float))}


def num(x):
    return isinstance(x, (int, float))


def load_hs861m():
    wb = openpyxl.load_workbook(os.path.join(RAW, "HS861M_2010.xlsx"), read_only=True, data_only=True)
    mon = {}
    for r in wb["Monthly-States"].iter_rows(min_row=4, values_only=True):
        if num(r[0]) and num(r[1]):
            mon[(r[2], int(r[0]), int(r[1]))] = dict(status=r[3], rev=r[4], mwh=r[5], cust=r[6], price=r[7])
    for r in wb["US-YTD"].iter_rows(min_row=4, values_only=True):
        if num(r[0]) and num(r[1]):
            mon[("US", int(r[0]), int(r[1]))] = dict(status=r[2], rev=r[3], mwh=r[4], cust=r[5], price=r[6])
    ytd = {}
    for r in wb["State-YTD-States"].iter_rows(min_row=4, values_only=True):
        if num(r[0]):
            ytd[(r[1], int(r[0]))] = dict(status=r[2], rev=r[3], mwh=r[4], cust=r[5], price=r[6])
    return mon, ytd


def load_hs861():
    wb = openpyxl.load_workbook(os.path.join(RAW, "HS861_2010.xlsx"), read_only=True, data_only=True)
    d = {}
    for r in wb["Total Electric Industry"].iter_rows(min_row=4, values_only=True):
        if num(r[0]):
            d[(r[1], int(r[0]))] = dict(rev=r[2], mwh=r[3], cust=r[4], price=r[5])
    return d


def load_pjm_share():
    """Доля бытовых потребителей штата у поставщиков с кодом балансирующей зоны PJM (EIA-861 2024).
    Считаются части A (bundled) и C (delivery) — так сумма равна числу потребителей в табл. 5A;
    часть B (energy-only) пропущена, чтобы не считать потребителя дважды."""
    wb = openpyxl.load_workbook(os.path.join(RAW, "Sales_Ult_Cust_2024.xlsx"), read_only=True, data_only=True)
    tot, pjm = defaultdict(float), defaultdict(float)
    for r in wb["States"].iter_rows(min_row=4, values_only=True):
        st, stype, ba, cust = r[6], r[4], r[8], r[11]
        if stype == "Energy" or not num(cust):
            continue
        tot[st] += cust
        if ba == "PJM":
            pjm[st] += cust
    return {s: (100 * pjm[s] / tot[s], pjm[s], tot[s]) for s in tot}


def load_fred():
    out = {}
    for s in FRED_IDS:
        d = {}
        for row in csv.reader(open(os.path.join(RAW, f"fred_{s}.csv"))):
            if row[0][:2] in ("19", "20"):
                d[row[0][:7]] = float(row[1]) if row[1] not in ("", ".") else None
        out[s] = d
    return out


def pct(a, b):
    return 100 * (a / b - 1)


def r1(x):
    return round(x + 0.0, 1)


def r2(x):
    return round(x + 0.0, 2)


# ---------------------------------------------------------------- типографика для текста
PROTECT = ["5.6.A", "5A", "CUSR0000SEHF01", "CUUR0000SEHF01", "CPIAUCNS", "EIA-861M", "EIA-861", "HS861M", "HS861"]
NBSP = " "


def ru(t):
    parts = re.split(r"(https?://\S+)", t)
    out = []
    for part in parts:
        if part.startswith("http"):
            out.append(part)
            continue
        for i, tok in enumerate(PROTECT):
            part = part.replace(tok, "@@%d@@" % i)
        dates = []

        def keep(m):
            dates.append(m.group(0))
            return "##%d##" % (len(dates) - 1)
        part = re.sub(r"\b\d{1,2}\.\d{2}\.\d{4}\b", keep, part)
        part = re.sub(r"(?<=\d)\.(?=\d)", ",", part)                       # десятичная запятая
        part = re.sub(r"(?<=\d) (?=(¢|кВт·ч|п\.\s?п\.|млн|млрд|раз|месяц))", NBSP, part)
        part = part.replace("п. п.", "п." + NBSP + "п.")
        part = re.sub(r"(?<![\w\d])-(?=\d)", "−", part)                      # минус
        for i, d in enumerate(dates):
            part = part.replace("##%d##" % i, d)
        for i, tok in enumerate(PROTECT):
            part = part.replace("@@%d@@" % i, tok)
        out.append(part)
    return "".join(out)


def g(x, nd=0):
    """Число с неразрывным пробелом между разрядами (1 047) — для текста реестра."""
    return f"{x:,.{nd}f}".replace(",", NBSP)


def main():
    fetch("--refresh" in sys.argv)
    epm_txt = page_text("epm_index.html")
    m = re.search(r"Data for (\w+ \d{4}) Release Date: (\w+ \d+, \d{4}) Next Release Date: (\w+ \d+, \d{4})", epm_txt)
    epm_month, epm_rel, epm_next = m.groups()
    assert epm_month == "July 2026", epm_month
    enduse = page_text("epm_update.html")
    q_enduse = "The largest percent increase was in Hawaii, up 25.4%, followed by Ohio, up 15.2%, and Maryland, up 14.8%."
    assert q_enduse in enduse
    q_proxy = re.search(r"we calculate average retail revenues per kWh as a proxy for retail rates and prices", enduse).group(0)
    pjm_txt = page_text("pjm_territory.html")
    assert "all or parts of Delaware, Illinois, Indiana, Kentucky, Maryland, Michigan, New Jersey, North Carolina, Ohio, Pennsylvania, Tennessee, Virginia, West Virginia and the District of Columbia" in pjm_txt
    b07, b08 = page_text("bls_cpi_08122026.htm"), page_text("bls_cpi_09112026.htm")
    assert "Over the last 12 months, the all items index increased 3.4 percent before seasonal adjustment" in b07
    assert "the electricity index rose 4.2 percent" in b07
    assert "the electricity index rose 3.8 percent" in b08
    assert "increased 3.4 percent before seasonal adjustment" in b08

    t56_title, t56 = load_t56a()
    t5a = load_t5a()
    mon, ytd = load_hs861m()
    ann = load_hs861()
    share = load_pjm_share()
    fred = load_fred()

    SRC_STATES = (f"EIA Electric Power Monthly, табл. 5.6.A (июль 2026, выпуск 24.09.2026) {U['t56a']} ; "
                  f"EIA ESRAP 2024, табл. 5A {U['t5a']} ; EIA-861M {U['hs861m']} ; "
                  f"EIA-861 2024, Sales_Ult_Cust (коды балансирующих зон) {U['f861']} ; PJM {U['pjm']}")
    cols = ["state", "state_ru", "pjm_coverage", "month", "res_price_cents", "res_price_cents_prev_year",
            "res_change_pct", "all_price_cents", "all_change_pct", "avg_kwh_month", "avg_bill_usd",
            "bill_change_usd_month", "source",
            # дополнительные колонки
            "pjm_share_res_customers_pct", "kwh_year", "all_price_cents_prev_year",
            "res_price_cents_2y_ago", "res_change_2y_pct", "bill_change_2y_usd_month",
            "avg_bill_2024_eia_usd", "july_bill_actual_2025_usd", "july_bill_actual_2026_usd"]
    rows, S = [], {}
    for name, ab, rus, cov, expl in STATES + [US]:
        key = "U.S. Total" if ab == "US" else name
        t, k = t56[key], t5a[key]
        m24 = mon[(ab, 2024, 7)]
        p24 = r2(m24["price"])
        kwh = k["kwh"]
        if ab == "US":
            covtxt, sh = f"{cov}: {expl}", ""
        else:
            sh_v = share[ab][0]
            sh_s = "100" if sh_v >= 99.95 else f"{sh_v:.1f}" if sh_v >= 90 or sh_v < 10 else f"{sh_v:.0f}"
            covtxt = ru(f"{cov}: " + expl.format(s=sh_s))
            sh = r1(sh_v)
        j25, j26 = mon[(ab, 2025, 7)], mon[(ab, 2026, 7)]
        row = dict(
            state=name, state_ru=rus, pjm_coverage=covtxt, month="2026-07",
            res_price_cents=t["res"], res_price_cents_prev_year=t["res_prev"],
            res_change_pct=r1(pct(t["res"], t["res_prev"])),
            all_price_cents=t["all"], all_change_pct=r1(pct(t["all"], t["all_prev"])),
            avg_kwh_month=r1(kwh), avg_bill_usd=r2(t["res"] * kwh / 100),
            bill_change_usd_month=r2((t["res"] - t["res_prev"]) * kwh / 100), source=SRC_STATES,
            pjm_share_res_customers_pct=sh, kwh_year=2024, all_price_cents_prev_year=t["all_prev"],
            res_price_cents_2y_ago=p24, res_change_2y_pct=r1(pct(t["res"], p24)),
            bill_change_2y_usd_month=r2((t["res"] - p24) * kwh / 100),
            avg_bill_2024_eia_usd=r2(k["bill"]),
            july_bill_actual_2025_usd=r2(j25["rev"] * 1000 / j25["cust"]),
            july_bill_actual_2026_usd=r2(j26["rev"] * 1000 / j26["cust"]),
        )
        # сверка: цена в табл. 5.6.A = цена в EIA-861M (обе округлены до сотых)
        assert abs(r2(j26["price"]) - t["res"]) < 0.006 and abs(r2(j25["price"]) - t["res_prev"]) < 0.006, (ab, j26["price"], t["res"])
        # сверка: счёт в табл. 5A = цена × потребление
        assert abs(k["price"] * k["kwh"] / 100 - k["bill"]) < 0.01
        rows.append(row)
        S[ab] = row
        # потребление 2025 (EIA-861M, предварительно) — для оговорки
        y25 = ytd[(ab, 2025)]
        row["_kwh25"] = y25["mwh"] * 1000 / y25["cust"] / 12
        row["_com"] = r1(pct(t["com"], t["com_prev"]))
        row["_jkwh25"], row["_jkwh26"] = j25["mwh"] * 1000 / j25["cust"], j26["mwh"] * 1000 / j26["cust"]
    with open(os.path.join(OUT, "03-states.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)

    # ------------------------------------------------ годовой ряд 2019–2025
    SRC_A = f"EIA-861 (HS861 2010-, окончательные) {U['hs861']}"
    SRC_M = f"EIA-861M (HS861M 2010-, сумма 12 месяцев, предварительно) {U['hs861m']}"
    arows, A = [], {}
    for name, ab, rus, *_ in STATES + [US]:
        prev = None
        for y in range(2019, 2026):
            if y <= 2024:
                p, status, src = ann[(ab, y)]["price"], "final", SRC_A
                assert abs(r2(ytd[(ab, y)]["price"]) - p) < 0.006          # два файла EIA согласованы
            else:
                p, status, src = r2(ytd[(ab, y)]["price"]), "preliminary", SRC_M
            A[(ab, y)] = p
            arows.append(dict(state=name, year=y, res_price_cents=p, state_ru=rus,
                              yoy_change_pct="" if prev is None else r1(pct(p, prev)),
                              data_status=status, source=src))
            prev = p
    with open(os.path.join(OUT, "03-states-annual.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["state", "year", "res_price_cents", "state_ru", "yoy_change_pct", "data_status", "source"])
        w.writeheader()
        w.writerows(arows)

    # ------------------------------------------------ ИПЦ
    cpi, el, elsa = fred["CPIAUCNS"], fred["CUUR0000SEHF01"], fred["CUSR0000SEHF01"]
    cpi_y = pct(cpi["2026-07"], cpi["2025-07"])
    el_y = pct(el["2026-07"], el["2025-07"])
    el_y8 = pct(el["2026-08"], el["2025-08"])
    cpi_y8 = pct(cpi["2026-08"], cpi["2025-08"])
    cpi_19 = pct(cpi["2026-07"], cpi["2019-07"])
    el_19 = pct(el["2026-07"], el["2019-07"])
    eia_19 = pct(t56["U.S. Total"]["res"], r2(mon[("US", 2019, 7)]["price"]))
    elsa_mm = pct(elsa["2026-08"], elsa["2026-07"])
    missing = [k for k, v in cpi.items() if v is None]

    # ------------------------------------------------ реестр
    C = "id,block,claim,plan_value,value,unit,date,tag,status,source,source_type,quote,note".split(",")
    R = []

    def r(*a):
        R.append(dict(zip(C, a)))

    P = [s[1] for s in STATES]
    FULL = ["DE", "DC", "MD", "NJ", "OH", "PA", "WV", "VA"]
    PART = ["IL", "IN", "KY", "MI", "NC", "TN"]
    RU = {s[1]: s[2] for s in STATES + [US]}

    def lst(key, codes, nd=1, money=False, sort=True):
        items = [(S[c][key], c) for c in codes]
        if sort:
            items.sort(key=lambda x: -x[0])
        fmt = (lambda v: ("+" if v >= 0 else "−") + "$" + f"{abs(v):.{nd}f}") if money else (lambda v: f"{v:+.{nd}f}")
        return "; ".join(f"{RU[c]} {fmt(v)}" for v, c in items)

    us = S["US"]
    E56 = U["t56a"]
    r("03-33", "03", "Самый свежий месяц в данных EIA о ценах по штатам на 29.09.2026 — июль 2026", "2026-07", "2026-07",
      "месяц данных", "2026-09-24", "ФАКТ", "проверено", U["epm"], "первоисточник",
      f"Data for {epm_month} Release Date: {epm_rel} Next Release Date: {epm_next}",
      "Electric Power Monthly, выпуск 24.09.2026; следующий — 23.10.2026 (данные за август 2026). Табл. 5.6.A, скачанная 28 и 29 сентября, совпадает побайтно. Июльские данные 2026 года предварительные (EIA-861M) и будут уточнены.")
    r("03-34", "03", "Средняя цена электричества для домохозяйств в США, июль 2026", "18.31", str(us["res_price_cents"]),
      "¢ за кВт·ч", "2026-07", "ФАКТ", "проверено", E56, "первоисточник",
      f"U.S. Total / Residential: July 2026 {us['res_price_cents']}; July 2025 {us['res_price_cents_prev_year']}",
      f"Годом раньше — {us['res_price_cents_prev_year']} ¢. По всем секторам — {us['all_price_cents']} против {us['all_price_cents_prev_year']} ¢. Июль — летний месяц, среднегодовая цена ниже: за 2025 год — {A[('US', 2025)]:.2f} ¢ (03-50).")
    r("03-35", "03", "Рост цены для домохозяйств в США за год", "4.9", str(us["res_change_pct"]), "% г/г", "2026-07",
      "ФАКТ", "проверено", E56, "данные",
      f"расчёт: {us['res_price_cents']} / {us['res_price_cents_prev_year']} = {us['res_price_cents'] / us['res_price_cents_prev_year']:.4f}",
      f"Номинальный рост, июль к июлю (цены взяты из таблицы с округлением до сотых; по точным выручке и объёму продаж EIA-861M — {pct(mon[('US', 2026, 7)]['price'], mon[('US', 2025, 7)]['price']):.2f}%). По всем секторам {us['all_change_pct']:+.1f}%. Инфляция (ИПЦ) за тот же период — {cpi_y:+.1f}% (03-52): свет для домохозяйств дорожал быстрее цен в целом.")
    oh, md = S["OH"], S["MD"]
    r("03-36", "03", "Огайо: рост средней цены по всем секторам за год", "15.2", str(oh["all_change_pct"]), "% г/г", "2026-07",
      "ФАКТ", "проверено", U["enduse"], "первоисточник", q_enduse,
      f"Все секторы: {oh['all_price_cents']:.2f} против {oh['all_price_cents_prev_year']:.2f} ¢ (табл. 5.6.A). Рост тянет коммерческий сектор: {t56['Ohio']['com']:.2f} против {t56['Ohio']['com_prev']:.2f} ¢ ({oh['_com']:+.1f}%); у домохозяйств {oh['res_change_pct']:+.1f}% (03-38). Первое место в США — Гавайи (+25.4%), не PJM.")
    r("03-37", "03", "Мэриленд: рост средней цены по всем секторам за год", "14.8", str(md["all_change_pct"]), "% г/г", "2026-07",
      "ФАКТ", "проверено", U["enduse"], "первоисточник", "followed by Ohio, up 15.2%, and Maryland, up 14.8%",
      f"Все секторы: {md['all_price_cents']:.2f} против {md['all_price_cents_prev_year']:.2f} ¢. Коммерческий сектор {md['_com']:+.1f}%, домохозяйства {md['res_change_pct']:+.1f}% (03-39).")
    r("03-38", "03", "Огайо: рост цены для домохозяйств за год", "11.9", str(oh["res_change_pct"]), "% г/г", "2026-07",
      "ФАКТ", "проверено", E56, "первоисточник",
      f"Ohio / Residential: July 2026 {oh['res_price_cents']}; July 2025 {oh['res_price_cents_prev_year']}",
      f"Для калькулятора: +${oh['bill_change_usd_month']:.2f} в месяц при потреблении {g(oh['avg_kwh_month'])} кВт·ч (среднее 2024). Огайо целиком в PJM.")
    r("03-39", "03", "Мэриленд: рост цены для домохозяйств за год", "13.7", str(md["res_change_pct"]), "% г/г", "2026-07",
      "ФАКТ", "проверено", E56, "первоисточник",
      f"Maryland / Residential: July 2026 {md['res_price_cents']}; July 2025 {md['res_price_cents_prev_year']}",
      f"Максимум среди штатов PJM. Для калькулятора: +${md['bill_change_usd_month']:.2f} в месяц при {g(md['avg_kwh_month'])} кВт·ч (среднее 2024). Мэриленд целиком в PJM.")
    lo = min(P, key=lambda c: S[c]["res_change_pct"])
    hi = max(P, key=lambda c: S[c]["res_change_pct"])
    assert lo == "NJ" and hi == "MD"
    r("03-40", "03", "Разброс роста цены для домохозяйств в штатах PJM за год", "-0.5…11.6",
      f"{S[lo]['res_change_pct']}…{S[hi]['res_change_pct']}", "% г/г", "2026-07", "ФАКТ", "исправлено", E56, "данные",
      f"Maryland / Residential: {md['res_price_cents']} vs {md['res_price_cents_prev_year']}; New Jersey / Residential: {S['NJ']['res_price_cents']} vs {S['NJ']['res_price_cents_prev_year']}",
      "Нижняя граница верна (Нью-Джерси), верхняя — Мэриленд (+13.7%), а не Иллинойс (+11.6% — только пятое место). Все штаты: " + lst("res_change_pct", P)
      + ". Штаты целиком или почти целиком в PJM — тот же диапазон: от Нью-Джерси до Мэриленда.")
    nj = S["NJ"]
    r("03-41", "03", "Нью-Джерси: цена для домохозяйств за год не выросла — скачок пришёлся на лето 2025", "-0.5",
      str(nj["res_change_pct"]), "% г/г", "2026-07", "ФАКТ", "проверено", U["hs861m"], "данные",
      f"NJ / Residential Price: July 2024 {nj['res_price_cents_2y_ago']}; July 2025 {nj['res_price_cents_prev_year']}; July 2026 {nj['res_price_cents']}",
      f"Эффект базы: июль 2024 → июль 2025 — {pct(nj['res_price_cents_prev_year'], nj['res_price_cents_2y_ago']):+.1f}%, за два года {nj['res_change_2y_pct']:+.1f}%. Поставочный год PJM начинается 1 июня, поэтому июль 2025 уже содержит цену мощности 2025/26 (03-03). Сравнение «июль к июлю» ловит только второй шаг; для калькулятора полезно показать и двухлетнее сравнение.")
    r("03-42", "03", "Рост цены для домохозяйств за два года (июль 2024 → июль 2026) в штатах PJM", "",
      f"{min(S[c]['res_change_2y_pct'] for c in P)}…{max(S[c]['res_change_2y_pct'] for c in P)}", "% за 2 года", "2026-07",
      "ФАКТ", "проверено", U["hs861m"], "данные",
      f"DC / Residential Price: July 2024 {S['DC']['res_price_cents_2y_ago']}; July 2026 {S['DC']['res_price_cents']}",
      "Июль 2024 — последний июль до поставочного года PJM 2025/26. По штатам: " + lst("res_change_2y_pct", P)
      + f"; США {us['res_change_2y_pct']:+.1f}. Цена мощности — лишь одна из причин роста: в те же годы менялись тарифы на передачу, распределение и стоимость топлива; причинность этими данными не доказывается.")
    r("03-43", "03", "Калькулятор: сколько долларов в месяц добавил рост цены за год (июль 2026 к июлю 2025)", "",
      f"{min(S[c]['bill_change_usd_month'] for c in P)}…{max(S[c]['bill_change_usd_month'] for c in P)}", "$ в месяц", "2026-07",
      "ОЦЕНКА", "проверено", E56 + " ; " + U["t5a"], "данные",
      f"расчёт для Мэриленда: ({md['res_price_cents']} − {md['res_price_cents_prev_year']}) × {md['avg_kwh_month']:.1f} кВт·ч / 100 = {md['bill_change_usd_month']:.2f}",
      "Допущение: потребление домохозяйства неизменно и равно среднемесячному за 2024 год (EIA, табл. 5A; данные за 2025 выйдут в октябре 2026); к нему применена разница средних цен июля 2026 и июля 2025. Средняя цена EIA — выручка/кВт·ч, поэтому Δцены × кВт·ч = Δсчёта при том же потреблении. "
      + lst("bill_change_usd_month", P, 2, True) + f"; США +${us['bill_change_usd_month']:.2f}.")
    r("03-44", "03", "Калькулятор: прибавка к месячному счёту за два года (июль 2026 к июлю 2024)", "",
      f"{min(S[c]['bill_change_2y_usd_month'] for c in P)}…{max(S[c]['bill_change_2y_usd_month'] for c in P)}", "$ в месяц", "2026-07",
      "ОЦЕНКА", "проверено", U["hs861m"] + " ; " + U["t5a"], "данные",
      f"расчёт для Огайо: ({oh['res_price_cents']} − {oh['res_price_cents_2y_ago']}) × {oh['avg_kwh_month']:.1f} / 100 = {oh['bill_change_2y_usd_month']:.2f}",
      "То же допущение (потребление 2024 года). По штатам: " + lst("bill_change_2y_usd_month", P, 2, True) + f"; США +${us['bill_change_2y_usd_month']:.2f}.")
    kw = sorted(P, key=lambda c: S[c]["avg_kwh_month"])
    dif25 = {c: pct(S[c]["_kwh25"], S[c]["avg_kwh_month"]) for c in P + ["US"]}
    r("03-45", "03", "Среднемесячное потребление электричества домохозяйством, 2024", "", f"{us['avg_kwh_month']:.1f}",
      "кВт·ч в месяц (США)", "2024", "ФАКТ", "проверено", U["t5a"], "первоисточник",
      f"U.S. Total / Average Monthly Consumption (kWh): {t5a['U.S. Total']['kwh']:.2f}",
      f"Штаты PJM: от {g(S[kw[0]]['avg_kwh_month'])} ({RU[kw[0]]}) до {g(S[kw[-1]]['avg_kwh_month'])} кВт·ч ({RU[kw[-1]]}); Огайо {g(oh['avg_kwh_month'])}, Пенсильвания {g(S['PA']['avg_kwh_month'])}, Вирджиния {g(S['VA']['avg_kwh_month'])}. Табл. 5A за 2025 год выйдет в октябре 2026. Сверка по EIA-861M за 2025 (предварительно): отличие от 2024 от {min(dif25[c] for c in P):+.1f}% ({RU[min(P, key=dif25.get)]}) до {max(dif25[c] for c in P):+.1f}% ({RU[max(P, key=dif25.get)]}), США {dif25['US']:+.1f}%: калькулятор на потреблении 2024 года скорее немного занижает прибавку.")
    bl = sorted(P, key=lambda c: S[c]["avg_bill_2024_eia_usd"])
    r("03-46", "03", "Средний месячный счёт домохозяйства за электричество, 2024", "", f"{us['avg_bill_2024_eia_usd']:.2f}",
      "$ в месяц (США)", "2024", "ФАКТ", "проверено", U["t5a"], "первоисточник",
      f"U.S. Total / Average Monthly Bill (Dollar and cents): {t5a['U.S. Total']['bill']:.2f}",
      f"Штаты PJM: от ${S[bl[0]]['avg_bill_2024_eia_usd']:.2f} ({RU[bl[0]]}) до ${S[bl[-1]]['avg_bill_2024_eia_usd']:.2f} ({RU[bl[-1]]}). В калькуляторе поле avg_bill_usd — оценка счёта по цене июля 2026 при потреблении 2024 года (США ${us['avg_bill_usd']:.2f}), а avg_bill_2024_eia_usd — фактическое среднее EIA за 2024.")
    full_s = "; ".join(f"{RU[c]} {share[c][0]:.1f}%" for c in FULL)
    r("03-47", "03", "PJM охватывает 13 штатов и округ Колумбия — одни целиком, другие частично", "13", "13",
      "штатов + DC", "2026-09-29", "ФАКТ", "проверено", U["pjm"], "первоисточник",
      "coordinates the movement of electricity through all or parts of Delaware, Illinois, Indiana, Kentucky, Maryland, Michigan, New Jersey, North Carolina, Ohio",
      "Доля бытовых потребителей штата у поставщиков с кодом балансирующей зоны PJM (EIA-861 2024, наш расчёт): " + full_s
      + ". Частично: " + "; ".join(f"{RU[c]} {share[c][0]:.1f}%" for c in PART)
      + ". Код зоны — на уровне компании, поэтому доли приблизительные.")
    il = S["IL"]
    r("03-48", "03", "Иллинойс входит в PJM только зоной ComEd (север штата с Чикаго)", "", f"{share['IL'][0]:.1f}",
      "% бытовых потребителей штата", "2024", "ФАКТ", "проверено", U["f861"], "данные",
      "Commonwealth Edison Co / IL / BA Code: PJM / Residential Customers: 2922223 bundled + 805452 delivery",
      f"Ameren Illinois (юг штата) — в MISO. Средняя цена EIA по Иллинойсу ({il['res_change_pct']:+.1f}% г/г) смешивает обе зоны; для PJM-части точнее тарифы ComEd.")
    r("03-49", "03", "Индиана, Кентукки, Мичиган, Северная Каролина и Теннесси — в PJM лишь меньшинство потребителей", "",
      f"{min(share[c][0] for c in PART if c != 'IL'):.1f}…{max(share[c][0] for c in PART if c != 'IL'):.1f}",
      "% бытовых потребителей штата", "2024", "ФАКТ", "проверено", U["f861"], "данные",
      f"Sales_Ult_Cust_2024 / KY: BA Code PJM {share['KY'][1]:.0f} of {share['KY'][2]:.0f} residential customers",
      "; ".join(f"{RU[c]} {share[c][0]:.1f}%" for c in PART if c != "IL")
      + f". В этих штатах средняя цена EIA описывает в основном энергосистему вне PJM: например, {S['NC']['res_change_pct']:+.1f}% в Северной Каролине — это прежде всего Duke Energy (вне RTO), {S['MI']['res_change_pct']:+.1f}% в Мичигане — DTE и Consumers (MISO). В калькуляторе такие штаты показывать с пометкой.")
    r("03-50", "03", "Годовая средняя цена для домохозяйств в США: 2019 → 2025", "", f"{A[('US', 2025)]:.2f}",
      "¢ за кВт·ч (2025)", "2025", "ФАКТ", "проверено", U["hs861m"], "данные",
      f"US-YTD / 2025 / Residential Price: {ytd[('US', 2025)]['price']:.4f} (Preliminary)",
      f"2019 — {A[('US', 2019)]:.2f} ¢, 2025 — {A[('US', 2025)]:.2f} ¢ ({pct(A[('US', 2025)], A[('US', 2019)]):+.1f}%). 2019–2024 — окончательные годовые данные EIA-861; 2025 — сумма 12 месяцев EIA-861M (предварительно, окончательные — в октябре 2026). Для 2019–2024 оба файла совпадают до сотых.")
    y25 = sorted(P, key=lambda c: -pct(A[(c, 2025)], A[(c, 2024)]))
    r("03-51", "03", "Когда начался рост: первая волна в 2022–2023 годах по всей стране, вторая в 2025-м — сильнее в штатах PJM", "",
      f"{pct(A[('US', 2022)], A[('US', 2021)]):.1f}", "% г/г (США, 2022)", "2025", "ФАКТ", "проверено", U["hs861"] + " ; " + U["hs861m"], "данные",
      f"US / Residential Price: 2021 {A[('US', 2021)]}; 2022 {A[('US', 2022)]}",
      f"2020–2021 — почти без роста (США {pct(A[('US', 2020)], A[('US', 2019)]):+.1f}% и {pct(A[('US', 2021)], A[('US', 2020)]):+.1f}%); 2022 — США {pct(A[('US', 2022)], A[('US', 2021)]):+.1f}% (в PJM: " + "; ".join(f"{RU[c]} {pct(A[(c, 2022)], A[(c, 2021)]):+.1f}" for c in ["IL", "PA", "VA", "OH", "MD"])
      + f"); 2023 — США {pct(A[('US', 2023)], A[('US', 2022)]):+.1f}%, в PJM: " + "; ".join(f"{RU[c]} {pct(A[(c, 2023)], A[(c, 2022)]):+.1f}" for c in ["DC", "MD", "DE", "PA", "OH"])
      + f"; 2024 — США {pct(A[('US', 2024)], A[('US', 2023)]):+.1f}%. 2025 — США {pct(A[('US', 2025)], A[('US', 2024)]):+.1f}%, в PJM: " + "; ".join(f"{RU[c]} {pct(A[(c, 2025)], A[(c, 2024)]):+.1f}" for c in y25[:6])
      + ". С 1 июня 2025 действует цена мощности PJM $269.92 за МВт-день (2025/26), в 9.3 раза выше прежней (03-03).")
    r("03-52", "03", "Инфляция в США (ИПЦ, все товары и услуги), июль 2026 к июлю 2025", "", f"{cpi_y:.1f}", "% г/г",
      "2026-07", "ФАКТ", "проверено", U["bls07"], "первоисточник",
      "Over the last 12 months, the all items index increased 3.4 percent before seasonal adjustment.",
      f"CPI-U без сезонной корректировки: {cpi['2026-07']} против {cpi['2025-07']} (FRED CPIAUCNS, {pct(cpi['2026-07'], cpi['2025-07']):+.2f}%). Цена для домохозяйств по EIA ({us['res_change_pct']:+.1f}%) росла быстрее на {us['res_change_pct'] - cpi_y:.1f} п. п.; в Мэриленде ({md['res_change_pct']:+.1f}%) — в четыре раза быстрее инфляции.")
    r("03-53", "03", "Индекс цен на электричество в ИПЦ, июль 2026 к июлю 2025", "", f"{el_y:.1f}", "% г/г", "2026-07",
      "ФАКТ", "проверено", U["bls07"], "первоисточник", "the electricity index rose 4.2 percent",
      f"CUUR0000SEHF01: {el['2026-07']} против {el['2025-07']}. BLS и EIA меряют разное: BLS — цену сопоставимой услуги для городских потребителей, EIA — среднюю выручку за кВт·ч у всех домохозяйств ({us['res_change_pct']:+.1f}%).")
    r("03-54", "03", "Индекс цен на электричество в ИПЦ (CUSR0000SEHF01), последний месяц — август 2026", "",
      f"{elsa['2026-08']}", "индекс, 1982–84 = 100, с сезонной корректировкой", "2026-08", "ФАКТ", "проверено",
      U["fredser"].format("CUSR0000SEHF01"), "данные", f"CUSR0000SEHF01, 2026-08: {elsa['2026-08']}; 2026-07: {elsa['2026-07']}",
      f"За месяц {elsa_mm:+.1f}%. Без сезонной корректировки (CUUR0000SEHF01) за год {el_y8:+.1f}%, ИПЦ в целом {cpi_y8:+.1f}% (BLS, выпуск 11.09.2026, те же значения). В рядах FRED нет значения за октябрь 2025 — на сравнения год к году за июль и август это не влияет.")
    r("03-55", "03", "С июля 2019 по июль 2026 электричество в ИПЦ подорожало заметно сильнее, чем всё остальное", "",
      f"{el_19:.1f}", "% (электричество в ИПЦ, июль 2019 → июль 2026)", "2026-07", "ФАКТ", "проверено",
      U["fredser"].format("CUUR0000SEHF01"), "данные", f"CUUR0000SEHF01: 2019-07 {el['2019-07']}; 2026-07 {el['2026-07']}",
      f"ИПЦ в целом за тот же период {cpi_19:+.1f}% (CPIAUCNS {cpi['2019-07']} → {cpi['2026-07']}). Цена EIA для домохозяйств США, июль к июлю: {r2(mon[('US', 2019, 7)]['price'])} → {us['res_price_cents']} ¢ ({eia_19:+.1f}%).")
    jb = ["NJ", "MD", "OH", "PA", "VA", "IL", "DC"]
    r("03-56", "03", "Фактический июльский счёт зависит от погоды: в Нью-Джерси он снизился при той же цене", "",
      f"{nj['july_bill_actual_2026_usd']:.2f}", "$ (средний счёт за июль 2026, NJ)", "2026-07", "ФАКТ", "проверено",
      U["hs861m"], "данные",
      f"NJ, July 2026 / Residential: revenue {mon[('NJ', 2026, 7)]['rev']:.0f} thousand $, customers {mon[('NJ', 2026, 7)]['cust']:.0f}",
      "Выручка / число потребителей, июль 2025 → июль 2026: " + "; ".join(f"{RU[c]} ${S[c]['july_bill_actual_2025_usd']:.0f} → ${S[c]['july_bill_actual_2026_usd']:.0f}" for c in jb)
      + f". В Нью-Джерси потребление в июле упало с {g(nj['_jkwh25'])} до {g(nj['_jkwh26'])} кВт·ч. Июльский счёт выше среднемесячного; калькулятор на годовом потреблении занижает летний счёт, но корректно выделяет вклад цены.")
    r("03-57", "03", "Средняя цена EIA — выручка за кВт·ч, а не тариф", "", "", "", "2026-07", "ФАКТ", "проверено",
      U["enduse"], "первоисточник", q_proxy,
      "Включает все начисления в счёте (энергия, мощность, сети, постоянные сборы), поэтому годится для калькулятора счёта; отдельные тарифы и надбавку за мощность PJM из неё не выделить. Для частично входящих в PJM штатов цена — среднее по всему штату (03-48, 03-49).")
    alls = sorted(P, key=lambda c: S[c]["all_change_pct"])
    r("03-58", "03", "Разброс роста цены по всем секторам в штатах PJM за год", "",
      f"{S[alls[0]]['all_change_pct']}…{S[alls[-1]]['all_change_pct']}", "% г/г", "2026-07", "ФАКТ", "проверено", E56, "данные",
      f"Ohio / All Sectors: {oh['all_price_cents']:.2f} vs {oh['all_price_cents_prev_year']:.2f}; Indiana / All Sectors: {S['IN']['all_price_cents']:.2f} vs {S['IN']['all_price_cents_prev_year']:.2f}",
      "Именно по всем секторам считает рейтинг штатов обзор EIA Electricity Monthly Update (03-36). " + lst("all_change_pct", P) + f"; США {us['all_change_pct']:+.1f}. В тексте указывать, о каком охвате речь.")
    va = S["VA"]
    r("03-59", "03", "Вирджиния — столица дата-центров: цена для домохозяйств за год", "", str(va["res_change_pct"]), "% г/г",
      "2026-07", "ФАКТ", "проверено", E56, "первоисточник",
      f"Virginia / Residential: July 2026 {va['res_price_cents']}; July 2025 {va['res_price_cents_prev_year']}",
      f"За два года {va['res_change_2y_pct']:+.1f}% (июль 2024 — {va['res_price_cents_2y_ago']} ¢). Калькулятор: +${va['bill_change_usd_month']:.2f} в месяц за год и +${va['bill_change_2y_usd_month']:.2f} за два года при {g(va['avg_kwh_month'])} кВт·ч. Коммерческий сектор {va['_com']:+.1f}%. Вирджиния в PJM почти целиком ({share['VA'][0]:.1f}%).")

    for row in R:
        row["claim"], row["note"] = ru(row["claim"]), ru(row["note"])
        for k in ("value", "plan_value"):
            assert " " not in row[k] and "," not in row[k], (row["id"], k, row[k])
        n = len(row["quote"].split())
        assert n <= 25, (row["id"], n, row["quote"])
    with open(os.path.join(OUT, "claims-03a.csv"), "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=C, quoting=csv.QUOTE_MINIMAL)
        w.writeheader()
        w.writerows(R)
    print("states:", len(rows), "annual:", len(arows), "claims:", len(R), R[0]["id"], "…", R[-1]["id"], "missing CPI:", missing)


if __name__ == "__main__":
    main()
